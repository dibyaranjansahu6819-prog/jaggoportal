from datetime import timedelta
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from django.utils import timezone

from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import UserProfile
from courses.models import Course, Subject
from teachers.models import Teacher

from .models import DailySchoolStatus, VolunteerAssignment


class Admin2ExcelExportTests(TestCase):

    def setUp(self):
        self.client = APIClient()

        self.course = Course.objects.create(
            name="Science"
        )

        self.subject = Subject.objects.create(
            name="Mathematics"
        )

        self.admin2_user = User.objects.create_user(
            username="excel_admin2",
            email="excel_admin2@example.com",
            password="AdminPassword123",
        )

        UserProfile.objects.create(
            user=self.admin2_user,
            role="ADMIN2",
            is_active=True,
        )

        self.volunteer_user = User.objects.create_user(
            username="excel_volunteer",
            email="excel_volunteer@example.com",
            password="VolunteerPassword123",
        )

        UserProfile.objects.create(
            user=self.volunteer_user,
            role="TEACHER_VOLUNTEER",
            is_active=True,
        )

        self.volunteer = Teacher.objects.create(
            auth_user=self.volunteer_user,
            user_id="EXCM001",
            name="Excel Volunteer",
            course=self.course,
            joining_year=2026,
            subject=self.subject,
            email="excel_volunteer@example.com",
            whatsapp_number="9876543210",
            free_days=[
                "Monday",
                "Tuesday",
                "Wednesday",
                "Thursday",
                "Friday",
                "Saturday",
            ],
            password="",
        )

        DailySchoolStatus.objects.create(
            date=timezone.localdate(),
            status="REGULAR_CLASS",
            created_by=self.admin2_user,
        )

        self.client.force_authenticate(
            user=self.admin2_user
        )

    # =========================================================
    # HELPERS
    # =========================================================

    def create_assignment(
        self,
        assignment_date=None,
        assigned_class="ClassA",
        task="TEACHING",
        instruction="Teach mathematics.",
        email_status="SENT",
    ):
        if assignment_date is None:
            assignment_date = timezone.localdate()

        return VolunteerAssignment.objects.create(
            volunteer=self.volunteer,
            assignment_date=assignment_date,
            assigned_class=assigned_class,
            task=task,
            instruction=instruction,
            email_status=email_status,
            created_by=self.admin2_user,
        )

    def get_worksheet(self, response):
        workbook = load_workbook(
            filename=BytesIO(response.content)
        )

        return workbook.active

    def get_table_rows(self, worksheet):
        """
        Current production Excel layout:

        Row 1 -> Title
        Row 2 -> Date
        Row 3 -> Status / report information
        Row 4 -> additional report information
        Row 5 -> table headers
        Row 6+ -> assignment rows

        The production workbook freezes the sheet at A6.
        """

        rows = list(
            worksheet.iter_rows(
                values_only=True
            )
        )

        # Header row is row 5 (zero-based index 4).
        self.assertGreaterEqual(
            len(rows),
            5,
        )

        headers = rows[4]

        data_rows = [
            row
            for row in rows[5:]
            if any(
                value is not None
                for value in row
            )
        ]

        return headers, data_rows

    # =========================================================
    # BASIC EXPORT
    # =========================================================

    def test_admin2_can_export_excel(self):
        self.create_assignment()

        response = self.client.get(
            "/api/admin2/assignments/export-excel/"
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        self.assertIn(
            "attachment;",
            response["Content-Disposition"],
        )

        self.assertIn(
            ".xlsx",
            response["Content-Disposition"],
        )

    # =========================================================
    # VALID XLSX
    # =========================================================

    def test_exported_file_is_valid_xlsx(self):
        self.create_assignment()

        response = self.client.get(
            "/api/admin2/assignments/export-excel/"
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        workbook = load_workbook(
            filename=BytesIO(response.content)
        )

        self.assertGreaterEqual(
            len(workbook.sheetnames),
            1,
        )

        self.assertIsNotNone(
            workbook.active
        )

    # =========================================================
    # HEADERS
    # =========================================================

    def test_excel_contains_expected_headers(self):
        self.create_assignment()

        response = self.client.get(
            "/api/admin2/assignments/export-excel/"
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        worksheet = self.get_worksheet(
            response
        )

        headers, data_rows = self.get_table_rows(
            worksheet
        )

        expected_headers = [
            "Date",
            "Class",
            "Volunteer Name",
            "Volunteer ID",
            "Subject",
            "Role",
            "Instruction",
            "Volunteer Email",
            "Email Status",
            "Email Sent At",
        ]

        self.assertEqual(
            len(headers),
            len(expected_headers),
        )

        for header in expected_headers:
            self.assertIn(
                header,
                headers,
            )

        self.assertEqual(
            len(data_rows),
            1,
        )

    # =========================================================
    # ASSIGNMENT DATA
    # =========================================================

    def test_excel_contains_assignment_data(self):
        assignment = self.create_assignment(
            assigned_class="ClassB",
            task="CHECKING",
            instruction="Check completed homework.",
        )

        response = self.client.get(
            "/api/admin2/assignments/export-excel/"
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        worksheet = self.get_worksheet(
            response
        )

        headers, data_rows = self.get_table_rows(
            worksheet
        )

        self.assertEqual(
            len(data_rows),
            1,
        )

        data_row = data_rows[0]

        date_index = headers.index("Date")
        class_index = headers.index("Class")
        name_index = headers.index("Volunteer Name")
        user_id_index = headers.index("Volunteer ID")
        subject_index = headers.index("Subject")
        role_index = headers.index("Role")
        instruction_index = headers.index("Instruction")
        email_index = headers.index("Volunteer Email")
        email_status_index = headers.index("Email Status")

        self.assertEqual(
            str(data_row[date_index]),
            str(assignment.assignment_date),
        )

        self.assertEqual(
            data_row[class_index],
            "Class B",
        )

        self.assertEqual(
            data_row[name_index],
            "Excel Volunteer",
        )

        self.assertEqual(
            data_row[user_id_index],
            "EXCM001",
        )

        self.assertEqual(
            data_row[subject_index],
            "Mathematics",
        )

        self.assertEqual(
            data_row[role_index],
            "Checking",
        )

        self.assertEqual(
            data_row[instruction_index],
            "Check completed homework.",
        )

        self.assertEqual(
            data_row[email_index],
            "excel_volunteer@example.com",
        )

        self.assertEqual(
            data_row[email_status_index],
            "SENT",
        )

    # =========================================================
    # DATE FILTER
    # =========================================================

    def test_excel_export_accepts_date_filter(self):
        today = timezone.localdate()

        previous_date = (
            today - timedelta(days=2)
        )

        self.create_assignment(
            assignment_date=today,
            assigned_class="ClassA",
            task="TEACHING",
        )

        self.create_assignment(
            assignment_date=previous_date,
            assigned_class="ClassC",
            task="INVIGILATOR",
        )

        response = self.client.get(
            "/api/admin2/assignments/export-excel/",
            {
                "date": previous_date.isoformat(),
            },
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        worksheet = self.get_worksheet(
            response
        )

        headers, data_rows = self.get_table_rows(
            worksheet
        )

        self.assertEqual(
            len(data_rows),
            1,
        )

        data_row = data_rows[0]

        date_index = headers.index("Date")
        class_index = headers.index("Class")
        role_index = headers.index("Role")
        name_index = headers.index("Volunteer Name")

        self.assertEqual(
            str(data_row[date_index]),
            str(previous_date),
        )

        self.assertEqual(
            data_row[class_index],
            "Class C",
        )

        self.assertEqual(
            data_row[role_index],
            "Invigilator",
        )

        self.assertEqual(
            data_row[name_index],
            self.volunteer.name,
        )

    # =========================================================
    # EMPTY EXPORT
    # =========================================================

    def test_excel_export_handles_no_assignments(self):
        response = self.client.get(
            "/api/admin2/assignments/export-excel/"
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        worksheet = self.get_worksheet(
            response
        )

        rows = list(
            worksheet.iter_rows(
                values_only=True
            )
        )

        # Empty report still contains the report
        # header and table header.
        self.assertGreaterEqual(
            len(rows),
            5,
        )

        headers = rows[4]

        expected_headers = [
            "Date",
            "Class",
            "Volunteer Name",
            "Volunteer ID",
            "Subject",
            "Role",
            "Instruction",
            "Volunteer Email",
            "Email Status",
            "Email Sent At",
        ]

        for header in expected_headers:
            self.assertIn(
                header,
                headers,
            )

        data_rows = [
            row
            for row in rows[5:]
            if any(
                value is not None
                for value in row
            )
        ]

        self.assertEqual(
            len(data_rows),
            0,
        )

    # =========================================================
    # UNAUTHENTICATED
    # =========================================================

    def test_unauthenticated_user_cannot_export_excel(self):
        self.client.force_authenticate(
            user=None
        )

        response = self.client.get(
            "/api/admin2/assignments/export-excel/"
        )

        self.assertEqual(
            response.status_code,
            401,
        )

    # =========================================================
    # ADMIN1 CANNOT EXPORT
    # =========================================================

    def test_admin1_cannot_export_excel(self):
        admin1_user = User.objects.create_user(
            username="excel_admin1",
            email="excel_admin1@example.com",
            password="AdminPassword123",
        )

        UserProfile.objects.create(
            user=admin1_user,
            role="ADMIN1",
            is_active=True,
        )

        self.client.force_authenticate(
            user=admin1_user
        )

        response = self.client.get(
            "/api/admin2/assignments/export-excel/"
        )

        self.assertEqual(
            response.status_code,
            403,
        )

    # =========================================================
    # DATE FILTERING
    # =========================================================

    def test_excel_contains_only_selected_date_assignments(self):
        today = timezone.localdate()

        another_date = (
            today - timedelta(days=5)
        )

        self.create_assignment(
            assignment_date=today,
            assigned_class="ClassA",
            task="TEACHING",
        )

        VolunteerAssignment.objects.create(
            volunteer=self.volunteer,
            assignment_date=another_date,
            assigned_class="ClassD",
            task="CHECKING",
            instruction="Old assignment.",
            email_status="SENT",
            created_by=self.admin2_user,
        )

        response = self.client.get(
            "/api/admin2/assignments/export-excel/",
            {
                "date": today.isoformat(),
            },
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        worksheet = self.get_worksheet(
            response
        )

        headers, data_rows = self.get_table_rows(
            worksheet
        )

        self.assertEqual(
            len(data_rows),
            1,
        )

        data_row = data_rows[0]

        date_index = headers.index("Date")
        class_index = headers.index("Class")

        self.assertEqual(
            str(data_row[date_index]),
            str(today),
        )

        self.assertEqual(
            data_row[class_index],
            "Class A",
        )

    # =========================================================
    # FREEZE HEADER
    # =========================================================

    def test_excel_header_is_frozen(self):
        self.create_assignment()

        response = self.client.get(
            "/api/admin2/assignments/export-excel/"
        )

        self.assertEqual(
            response.status_code,
            200,
        )

        worksheet = self.get_worksheet(
            response
        )

        self.assertEqual(
            worksheet.freeze_panes,
            "A6",
        )