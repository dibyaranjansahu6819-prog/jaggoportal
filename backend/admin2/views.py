import io
import textwrap
from django.shortcuts import get_object_or_404
from django.core.mail import EmailMessage
from django.db import transaction
from django.db.models import Sum
from django.utils import timezone
from django.http import HttpResponse
from django.conf import settings

from PIL import Image, ImageDraw, ImageFont

from rest_framework import status
from rest_framework.parsers import (
    FormParser,
    JSONParser,
    MultiPartParser,
)
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.permissions import IsAdmin2
from attendance.models import (
    StudentAttendance,
    VolunteerAttendance,
)
from students.models import Student
from teachers.models import Teacher

from .models import (
    DailySchoolStatus,
    VolunteerAccountStatus,
    VolunteerAccessRequest,
    VolunteerAssignment,
    VolunteerCaution,
    VolunteerXPTransaction,
)
from .serializers import (
    AccessRequestSerializer,
    DailySchoolStatusSerializer,
    ManualXPSerializer,
    SendAssignmentSerializer,
    StudentSummarySerializer,
    VolunteerAccountStatusSerializer,
    VolunteerAssignmentSerializer,
    VolunteerCautionSerializer,
    VolunteerListSerializer,
    XPTransactionSerializer,
)
from .services import (
    get_assigned_absence_streak,
    get_current_status,
    get_volunteer_xp,
    send_playing_day_emails,
)


class Admin2BaseView(APIView):
    permission_classes = [
        IsAuthenticated,
        IsAdmin2,
    ]


class DailySchoolStatusView(Admin2BaseView):
    """Get or set today's school operating status."""

    def get(self, request):
        today = timezone.localdate()
        daily_status = DailySchoolStatus.objects.filter(
            date=today
        ).select_related("created_by").first()

        if not daily_status:
            return Response({
                "date": today,
                "status": None,
                "message": "Today's school status has not been selected yet.",
            })

        return Response(
            DailySchoolStatusSerializer(daily_status).data
        )

    @transaction.atomic
    def post(self, request):
        today = timezone.localdate()
        requested_status = str(
            request.data.get("status", "")
        ).upper().strip()

        valid_statuses = {
            "HOLIDAY",
            "PLAYING_DAY",
            "REGULAR_CLASS",
        }

        if requested_status not in valid_statuses:
            return Response(
                {
                    "detail": (
                        "Status must be HOLIDAY, PLAYING_DAY, "
                        "or REGULAR_CLASS."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        daily_status, created = DailySchoolStatus.objects.get_or_create(
            date=today,
            defaults={
                "status": requested_status,
                "created_by": request.user,
            },
        )

        # Changing to another status is allowed.
        old_status = daily_status.status
        daily_status.status = requested_status

        if requested_status != "PLAYING_DAY":
            daily_status.playing_day_email_status = "NOT_SENT"
            daily_status.playing_day_email_sent_at = None
            daily_status.playing_day_email_error = ""

        daily_status.save()

        result = None
        if requested_status == "PLAYING_DAY":
            # Send only when entering Playing Day or when the previous
            # notification failed/not sent. This avoids duplicate emails.
            if (
                created
                or old_status != "PLAYING_DAY"
                or daily_status.playing_day_email_status != "SENT"
            ):
                result = send_playing_day_emails(daily_status)

        response_data = DailySchoolStatusSerializer(
            daily_status
        ).data

        if result is not None:
            response_data["playing_day_email_result"] = result

        return Response(
            response_data,
            status=(
                status.HTTP_201_CREATED
                if created
                else status.HTTP_200_OK
            ),
        )


class Admin2DashboardView(Admin2BaseView):

    def get(self, request):
        today = timezone.localdate()

        total_students = Student.objects.count()

        present_students = (
            StudentAttendance.objects.filter(
                session__session_date=today,
                status="PRESENT",
            )
            .values("student_id")
            .distinct()
            .count()
        )

        assignments = VolunteerAssignment.objects.filter(
            assignment_date=today
        ).select_related(
            "volunteer",
            "created_by",
        )

        caution_count = VolunteerCaution.objects.filter(
            status="ACTIVE"
        ).count()

        pending_access = VolunteerAccessRequest.objects.filter(
            status="PENDING"
        ).count()

        return Response(
            {
                "date": today,
                "students": {
                    "total_registered": total_students,
                    "present_today": present_students,
                },
                "volunteers": {
                    "registered": Teacher.objects.count(),
                },
                "assignments_today": assignments.count(),
                "active_cautions": caution_count,
                "pending_access_requests": pending_access,
            }
        )


class TotalStudentsView(Admin2BaseView):

    def get(self, request):
        students = Student.objects.all().order_by(
            "student_class",
            "roll_no",
        )

        return Response(
            {
                "count": students.count(),
                "students": StudentSummarySerializer(
                    students,
                    many=True,
                ).data,
            }
        )


class PresentStudentsTodayView(Admin2BaseView):

    def get(self, request):
        today = timezone.localdate()

        student_ids = (
            StudentAttendance.objects.filter(
                session__session_date=today,
                status="PRESENT",
            )
            .values_list(
                "student_id",
                flat=True,
            )
            .distinct()
        )

        students = Student.objects.filter(
            id__in=student_ids
        ).order_by(
            "student_class",
            "roll_no",
        )

        return Response(
            {
                "date": today,
                "count": students.count(),
                "students": StudentSummarySerializer(
                    students,
                    many=True,
                ).data,
            }
        )


class RegisteredVolunteersView(Admin2BaseView):

    def get(self, request):
        volunteers = Teacher.objects.all().select_related(
            "course",
            "subject",
        ).order_by("name")

        return Response(
            {
                "count": volunteers.count(),
                "volunteers": VolunteerListSerializer(
                    volunteers,
                    many=True,
                ).data,
            }
        )


class TodayAssignmentsView(Admin2BaseView):

    def get(self, request):
        today = timezone.localdate()

        assignments = VolunteerAssignment.objects.filter(
            assignment_date=today
        ).select_related(
            "volunteer",
            "created_by",
        )

        return Response(
            {
                "date": today,
                "count": assignments.count(),
                "assignments": VolunteerAssignmentSerializer(
                    assignments,
                    many=True,
                ).data,
            }
        )


class AssignmentHistoryView(Admin2BaseView):

    def get(self, request):
        assignments = VolunteerAssignment.objects.all().select_related(
            "volunteer",
            "created_by",
        )

        return Response(
            {
                "count": assignments.count(),
                "assignments": VolunteerAssignmentSerializer(
                    assignments,
                    many=True,
                ).data,
            }
        )


class SendAssignmentView(Admin2BaseView):

    parser_classes = [
        MultiPartParser,
        FormParser,
        JSONParser,
    ]

    @transaction.atomic
    def post(self, request):
        serializer = SendAssignmentSerializer(
            data=request.data
        )

        serializer.is_valid(
            raise_exception=True
        )

        volunteer = serializer.validated_data[
            "volunteer"
        ]

        assignment = VolunteerAssignment(
            volunteer=volunteer,
            assignment_date=timezone.localdate(),
            assigned_class=serializer.validated_data[
                "assigned_class"
            ],
            task=serializer.validated_data[
                "task"
            ],
            instruction=serializer.validated_data.get(
                "instruction",
                "",
            ),
            attachment=serializer.validated_data.get(
                "attachment"
            ),
            homework_attachment=serializer.validated_data.get(
                "homework_attachment"
            ),
            sender_email="",
            created_by=request.user,
        )

        assignment.full_clean()

        class_name = (
            assignment.get_assigned_class_display()
        )

        task_name = (
            assignment.get_task_display()
        )

        subject = (
            "Jaago Team - Volunteer Assignment - "
            f"{class_name}"
        )

        body = (
            f"Hello {volunteer.name},\n\n"
            "You have been assigned the following task "
            f"for {assignment.assignment_date}.\n\n"
            f"Class: {class_name}\n"
            f"Task: {task_name}\n\n"
            "Instruction / Details:\n"
            f"{assignment.instruction or 'No additional instructions.'}\n\n"
            f"Teaching / Checking Module: {assignment.attachment.name.split('/')[-1] if assignment.attachment else 'Not provided'}\n"
            f"Homework: {assignment.homework_attachment.name.split('/')[-1] if assignment.homework_attachment else 'Not provided'}\n\n"
            "Please follow the assigned instructions.\n\n"
            "Regards, Jaago Team"
        )

        try:
            email = EmailMessage(
                subject=subject,
                body=body,
                to=[volunteer.email],
            )

            if assignment.attachment:
                email.attach_file(
                    assignment.attachment.path
                )

            if assignment.homework_attachment:
                email.attach_file(
                    assignment.homework_attachment.path
                )

            email.send(
                fail_silently=False
            )

            assignment.email_status = "SENT"
            assignment.email_sent_at = timezone.now()
            assignment.sender_email = (
                getattr(
                    email,
                    "from_email",
                    ""
                )
                or ""
            )

            assignment.save()

        except Exception as exc:
            assignment.email_status = "FAILED"
            assignment.email_error = str(exc)
            assignment.save()

            return Response(
                {
                    "detail": "Email could not be sent.",
                    "error": str(exc),
                    "assignment": VolunteerAssignmentSerializer(
                        assignment
                    ).data,
                },
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )

        return Response(
            {
                "message": "Assignment sent successfully.",
                "assignment": VolunteerAssignmentSerializer(
                    assignment
                ).data,
            },
            status=status.HTTP_201_CREATED,
        )


class VolunteerXPView(Admin2BaseView):

    def get(self, request, volunteer_id):
        volunteer = Teacher.objects.filter(
            id=volunteer_id
        ).first()

        if not volunteer:
            return Response(
                {
                    "detail": "Volunteer not found."
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        transactions = VolunteerXPTransaction.objects.filter(
            volunteer=volunteer
        )

        return Response(
            {
                "volunteer": volunteer.id,
                "user_id": volunteer.user_id,
                "name": volunteer.name,
                "xp": get_volunteer_xp(volunteer),
                "transactions": XPTransactionSerializer(
                    transactions,
                    many=True,
                ).data,
            }
        )


class ManualXPView(Admin2BaseView):

    @transaction.atomic
    def post(self, request):
        serializer = ManualXPSerializer(
            data=request.data
        )

        serializer.is_valid(
            raise_exception=True
        )

        volunteer = serializer.validated_data[
            "volunteer"
        ]

        points = serializer.validated_data[
            "points"
        ]

        reason = serializer.validated_data[
            "reason"
        ]

        source = (
            "ADMIN_INCREASE"
            if points > 0
            else "ADMIN_DECREASE"
        )

        transaction_record = (
            VolunteerXPTransaction.objects.create(
                volunteer=volunteer,
                points=points,
                source=source,
                reason=reason,
                created_by=request.user,
            )
        )

        return Response(
            {
                "message": "XP updated successfully.",
                "current_xp": get_volunteer_xp(
                    volunteer
                ),
                "transaction": XPTransactionSerializer(
                    transaction_record
                ).data,
            },
            status=status.HTTP_201_CREATED,
        )


class CautionListView(Admin2BaseView):

    def get(self, request):
        cautions = VolunteerCaution.objects.filter(
            status="ACTIVE"
        ).select_related(
            "volunteer"
        )

        return Response(
            {
                "count": cautions.count(),
                "cautions": VolunteerCautionSerializer(
                    cautions,
                    many=True,
                ).data,
            }
        )


class VolunteerStatusView(Admin2BaseView):

    def get(self, request, volunteer_id):
        volunteer = Teacher.objects.filter(
            id=volunteer_id
        ).first()

        if not volunteer:
            return Response(
                {
                    "detail": "Volunteer not found."
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        account_status = get_current_status(
            volunteer
        )

        return Response(
            VolunteerAccountStatusSerializer(
                account_status
            ).data
        )


class RemoveVolunteerView(Admin2BaseView):

    @transaction.atomic
    def post(self, request, volunteer_id):
        volunteer = Teacher.objects.filter(
            id=volunteer_id
        ).first()

        if not volunteer:
            return Response(
                {
                    "detail": "Volunteer not found."
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        reason = request.data.get(
            "reason",
            "Removed by Admin 2.",
        )

        account_status = get_current_status(
            volunteer
        )

        account_status.status = "REMOVED"
        account_status.removed_at = timezone.now()
        account_status.removed_by = request.user
        account_status.removal_reason = reason

        account_status.save()

        if volunteer.auth_user:
            volunteer.auth_user.is_active = False
            volunteer.auth_user.save(
                update_fields=["is_active"]
            )

        return Response(
            {
                "message": (
                    "Volunteer removed successfully."
                ),
                "status": "REMOVED",
                "volunteer": volunteer.user_id,
            }
        )


class AccessRequestListView(Admin2BaseView):

    def get(self, request):
        requests = VolunteerAccessRequest.objects.filter(
            status="PENDING"
        ).select_related(
            "volunteer"
        )

        return Response(
            {
                "count": requests.count(),
                "requests": AccessRequestSerializer(
                    requests,
                    many=True,
                ).data,
            }
        )


class DecideAccessRequestView(Admin2BaseView):

    @transaction.atomic
    def post(self, request, request_id):
        access_request = (
            VolunteerAccessRequest.objects
            .select_related("volunteer")
            .filter(id=request_id)
            .first()
        )

        if not access_request:
            return Response(
                {
                    "detail": (
                        "Access request not found."
                    )
                },
                status=status.HTTP_404_NOT_FOUND,
            )

        if access_request.status != "PENDING":
            return Response(
                {
                    "detail": (
                        "This access request has "
                        "already been decided."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        decision = str(
            request.data.get(
                "decision",
                ""
            )
        ).upper()

        if decision not in [
            "APPROVED",
            "DENIED",
        ]:
            return Response(
                {
                    "detail": (
                        "Decision must be APPROVED "
                        "or DENIED."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        note = request.data.get(
            "decision_note",
            "",
        )

        access_request.status = decision
        access_request.decided_at = timezone.now()
        access_request.decided_by = request.user
        access_request.decision_note = note

        access_request.save()

        if decision == "APPROVED":
            account_status = get_current_status(
                access_request.volunteer
            )

            account_status.status = "ACTIVE"
            account_status.removed_at = None
            account_status.removed_by = None

            account_status.save()

            if access_request.volunteer.auth_user:
                access_request.volunteer.auth_user.is_active = True
                access_request.volunteer.auth_user.save(
                    update_fields=["is_active"]
                )

        return Response(
            {
                "message": (
                    "Access request "
                    f"{decision.lower()}."
                ),
                "request": AccessRequestSerializer(
                    access_request
                ).data,
            }
        )


class DailyReportPNGView(Admin2BaseView):
    """Generate today's volunteer timetable as a PNG."""

    def get(self, request):
        today = timezone.localdate()
        daily_status = DailySchoolStatus.objects.filter(
            date=today
        ).first()

        assignments = list(
            VolunteerAssignment.objects.filter(
                assignment_date=today
            )
            .select_related(
                "volunteer",
                "volunteer__subject",
            )
            .order_by(
                "assigned_class",
                "volunteer__name",
            )
        )

        width = 1400
        margin = 60

        try:
            title_font = ImageFont.truetype("arialbd.ttf", 42)
            section_font = ImageFont.truetype("arialbd.ttf", 28)
            normal_font = ImageFont.truetype("arial.ttf", 22)
            small_font = ImageFont.truetype("arial.ttf", 18)
        except OSError:
            title_font = ImageFont.load_default()
            section_font = ImageFont.load_default()
            normal_font = ImageFont.load_default()
            small_font = ImageFont.load_default()

        rows = max(len(assignments), 1)
        height = 280 + rows * 100 + 100

        image = Image.new("RGB", (width, height), "white")
        draw = ImageDraw.Draw(image)

        y = 45
        draw.text((margin, y), "JAAGO PORTAL", fill="black", font=title_font)
        y += 60
        draw.text((margin, y), "Daily Volunteer Timetable", fill="black", font=section_font)
        y += 45
        draw.text((margin, y), f"Date: {today}", fill="black", font=normal_font)
        y += 38

        status_text = (
            daily_status.get_status_display()
            if daily_status
            else "Status not selected"
        )
        draw.text((margin, y), f"Status: {status_text}", fill="black", font=normal_font)
        y += 45
        draw.line((margin, y, width - margin, y), fill="black", width=3)
        y += 30

        columns = [
            ("Class", margin, 220),
            ("Volunteer", 300, 400),
            ("Subject", 710, 330),
            ("Role", 1050, 290),
        ]

        for title, x, column_width in columns:
            draw.rectangle((x, y, x + column_width, y + 55), outline="black", width=2)
            draw.text((x + 10, y + 15), title, fill="black", font=small_font)

        y += 55

        if assignments:
            for assignment in assignments:
                row_height = 100
                values = [
                    (assignment.get_assigned_class_display(), margin, 220),
                    (assignment.volunteer.name, 300, 400),
                    (assignment.volunteer.subject.name if assignment.volunteer.subject else "-", 710, 330),
                    (assignment.get_task_display(), 1050, 290),
                ]
                for value, x, column_width in values:
                    draw.rectangle((x, y, x + column_width, y + row_height), outline="black", width=2)
                    for line_index, line in enumerate(textwrap.wrap(str(value), width=max(10, int(column_width / 11)))[:3]):
                        draw.text((x + 10, y + 14 + line_index * 24), line, fill="black", font=small_font)
                y += row_height
        else:
            draw.rectangle((margin, y, width - margin, y + 100), outline="black", width=2)
            message = (
                "No volunteer assignments for today."
                if daily_status and daily_status.status == "REGULAR_CLASS"
                else "No volunteer timetable for this day."
            )
            draw.text((margin + 20, y + 30), message, fill="black", font=normal_font)
            y += 100

        y += 25
        draw.line((margin, y, width - margin, y), fill="black", width=2)
        y += 25
        draw.text((margin, y), "Generated by Jaago Portal", fill="black", font=small_font)

        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        buffer.seek(0)

        response = HttpResponse(buffer.getvalue(), content_type="image/png")
        response["Content-Disposition"] = (
            f'attachment; filename="jaago_daily_timetable_{today}.png"'
        )
        return response


class AssignmentExcelExportView(Admin2BaseView):
    """Export volunteer assignment timetable data as an Excel workbook.

    By default the export is for today's assignments. Admin2 can optionally
    provide ?date=YYYY-MM-DD to export the assignments for a specific date.
    The export intentionally contains timetable/assignment information only;
    attachment contents and XP/statistics are not included.
    """

    def get(self, request):
        selected_date_text = str(
            request.query_params.get("date", "")
        ).strip()

        if selected_date_text:
            try:
                selected_date = timezone.datetime.strptime(
                    selected_date_text,
                    "%Y-%m-%d",
                ).date()
            except ValueError:
                return Response(
                    {
                        "detail": "Date must be in YYYY-MM-DD format."
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )
        else:
            selected_date = timezone.localdate()

        assignments = list(
            VolunteerAssignment.objects.filter(
                assignment_date=selected_date
            )
            .select_related(
                "volunteer",
                "volunteer__subject",
                "created_by",
            )
            .order_by(
                "assigned_class",
                "volunteer__name",
            )
        )

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font
            from openpyxl.utils import get_column_letter
        except ImportError:
            return Response(
                {
                    "detail": (
                        "openpyxl is required for Excel export. "
                        "Install it with: pip install openpyxl"
                    )
                },
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Assignments"

        daily_status = DailySchoolStatus.objects.filter(
            date=selected_date
        ).first()
        status_text = (
            daily_status.get_status_display()
            if daily_status
            else "Status not selected"
        )

        worksheet.append(["JAAGO PORTAL - VOLUNTEER ASSIGNMENTS"])
        worksheet.append(["Date", selected_date.isoformat()])
        worksheet.append(["School Status", status_text])
        worksheet.append([])

        headers = [
            "Date",
            "Class",
            "Volunteer Name",
            "Volunteer ID",
            "Subject",
            "Role",
            "Instruction",
            "Volunteer Email",
            "Email Status",
            "Email Sent At",
        ]
        worksheet.append(headers)

        for cell in worksheet[1]:
            cell.font = Font(bold=True, size=16)

        for cell in worksheet[5]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        for assignment in assignments:
            worksheet.append([
                assignment.assignment_date.isoformat(),
                assignment.get_assigned_class_display(),
                assignment.volunteer.name,
                assignment.volunteer.user_id,
                (
                    assignment.volunteer.subject.name
                    if assignment.volunteer.subject
                    else "-"
                ),
                assignment.get_task_display(),
                assignment.instruction or "",
                assignment.volunteer.email or "",
                assignment.email_status,
                (
                    assignment.email_sent_at.isoformat()
                    if assignment.email_sent_at
                    else ""
                ),
            ])

        widths = {
            "A": 14,
            "B": 14,
            "C": 28,
            "D": 16,
            "E": 20,
            "F": 18,
            "G": 45,
            "H": 34,
            "I": 16,
            "J": 28,
        }
        for column, width in widths.items():
            worksheet.column_dimensions[column].width = width

        for row in worksheet.iter_rows(min_row=6):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

        worksheet.freeze_panes = "A6"
        worksheet.auto_filter.ref = (
            f"A5:J{max(5, worksheet.max_row)}"
        )

        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )
        response["Content-Disposition"] = (
            'attachment; '
            f'filename="jaago_assignments_{selected_date}.xlsx"'
        )
        return response


class RetryAssignmentEmailView(APIView):
    permission_classes = [IsAdmin2]

    def post(self, request, assignment_id):
        assignment = get_object_or_404(
            VolunteerAssignment.objects.select_related("volunteer"),
            id=assignment_id,
        )

        volunteer = assignment.volunteer

        if not volunteer.email:
            return Response(
                {"detail": "Volunteer does not have a registered email address."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if assignment.email_status != "FAILED":
            return Response(
                {
                    "detail": "Only failed assignments can be retried.",
                    "email_status": assignment.email_status,
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        subject = (
            f"Jaago Team - Volunteer Assignment - "
            f"{assignment.assigned_class}"
        )

        body = (
            f"Dear {volunteer.name},\n\n"
            f"You have been assigned the following work:\n\n"
            f"Date: {assignment.assignment_date}\n"
            f"Class: {assignment.assigned_class}\n"
            f"Task: {assignment.get_task_display()}\n"
            f"Instruction: {assignment.instruction or 'No additional instruction.'}\n\n"
        )

        if assignment.attachment:
            body += f"Teaching/Checking Module: {assignment.attachment.name}\n"

        if assignment.homework_attachment:
            body += f"Homework: {assignment.homework_attachment.name}\n"

        body += "\nRegards,\nJaago Team"

        email = EmailMessage(
            subject=subject,
            body=body,
            from_email=settings.DEFAULT_FROM_EMAIL,
            to=[volunteer.email],
        )

        if assignment.attachment:
            try:
                assignment.attachment.open("rb")
                email.attach(
                    os.path.basename(assignment.attachment.name),
                    assignment.attachment.read(),
                )
                assignment.attachment.close()
            except Exception:
                pass

        if assignment.homework_attachment:
            try:
                assignment.homework_attachment.open("rb")
                email.attach(
                    os.path.basename(assignment.homework_attachment.name),
                    assignment.homework_attachment.read(),
                )
                assignment.homework_attachment.close()
            except Exception:
                pass

        try:
            email.send(fail_silently=False)

            assignment.email_status = "SENT"
            assignment.email_sent_at = timezone.now()
            assignment.email_error = ""
            assignment.sender_email = settings.DEFAULT_FROM_EMAIL
            assignment.save(
                update_fields=[
                    "email_status",
                    "email_sent_at",
                    "email_error",
                    "sender_email",
                    "updated_at",
                ]
            )

            return Response(
                {
                    "message": "Assignment email sent successfully.",
                    "assignment": VolunteerAssignmentSerializer(
                        assignment,
                        context={"request": request},
                    ).data,
                },
                status=status.HTTP_200_OK,
            )

        except Exception as exc:
            assignment.email_status = "FAILED"
            assignment.email_error = str(exc)
            assignment.save(
                update_fields=[
                    "email_status",
                    "email_error",
                    "updated_at",
                ]
            )

            return Response(
                {
                    "detail": "Assignment email failed.",
                    "error": str(exc),
                    "assignment_id": assignment.id,
                },
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )